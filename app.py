import os
import re
import json
import secrets
import sqlite3
from datetime import datetime, timezone
import requests
from pathlib import Path
from flask import Flask, render_template, request, send_file, jsonify, session, redirect, url_for, abort
from werkzeug.exceptions import RequestEntityTooLarge
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from docx import Document
from PyPDF2 import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from content import BLOG_POSTS, BY_NAME, BY_SLUG, COUNTRIES, COUNTRY_GUIDANCE, FAQS, FEATURED, POSTS_BY_SLUG

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "jobaround-dev-secret-change-me")
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.getenv("COOKIE_SECURE", "0") == "1"

UPLOAD_DIR = "uploads"
OUTPUT_DIR = "outputs"
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:3b")

ALLOWED = {"pdf", "docx", "txt"}
DB_PATH = os.path.join("data", "jobaround.db")

def db():
    os.makedirs("data", exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = db()
    try:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS resumes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                country TEXT NOT NULL,
                job_title TEXT,
                preview TEXT,
                pdf_name TEXT,
                docx_name TEXT,
                ai_used INTEGER NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT NOT NULL,
                body TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
        """)
        conn.commit()
    finally:
        conn.close()

init_db()

def now():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]

def csrf_ok():
    token = request.form.get("csrf", "")
    return token and token == session.get("csrf")

def current_user():
    uid = session.get("user_id")
    if not uid:
        return None
    conn = db()
    try:
        row = conn.execute("SELECT id, name, email, created_at FROM users WHERE id = ?", (uid,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()

@app.context_processor
def inject_user():
    return {"user": current_user(), "csrf_token": csrf_token(), "year": 2026}

def extract_text(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if ext == ".txt":
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    raise ValueError("Unsupported file type")

def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def fallback_tailor(resume, country, job_title, job_description):
    """Useful fallback when Ollama is not reachable. It does not invent experience."""
    lines = [x.strip() for x in resume.splitlines() if x.strip()]
    keywords = []
    source = (job_title + " " + job_description).lower()
    for word in re.findall(r"[a-zA-Z][a-zA-Z0-9+#.-]{2,}", source):
        if word not in keywords:
            keywords.append(word)
    matched = []
    resume_lower = resume.lower()
    for word in keywords:
        if word in resume_lower and word not in matched:
            matched.append(word)
    summary = f"Professional with experience relevant to {job_title or 'the target role'}, tailored for the {country} job market."
    if matched:
        summary += " Relevant keywords already present in the resume include: " + ", ".join(matched[:12]) + "."
    return (
        f"TARGET ROLE\n{job_title or 'Not specified'}\n\n"
        f"TARGET COUNTRY\n{country}\n\n"
        f"PROFESSIONAL SUMMARY\n{summary}\n\n"
        f"ORIGINAL RESUME CONTENT\n{resume}\n\n"
        f"TAILORING NOTE\nThe automatic local-model service was not available, so the original facts were preserved and no new qualifications or experience were invented."
    )

def tailor_with_ollama(resume, country, job_title, job_description):
    guidance = COUNTRY_GUIDANCE.get(country, "Use the country's common professional resume conventions while avoiding unsupported assumptions.")
    prompt = f"""
You are Jobaround, a professional resume localization assistant.

TASK:
Rewrite the user's resume for a job application in {country}.
Target job title: {job_title or "Not specified"}.
Job description:
{job_description or "Not provided"}.

COUNTRY GUIDANCE:
{guidance}

STRICT RULES:
1. Never invent employers, dates, degrees, certifications, technologies, achievements, metrics, job titles or responsibilities.
2. Only rewrite, reorder, shorten or clarify information that is actually present in the source resume.
3. If a job description contains a keyword that is not supported by the resume, do not claim the user has that skill.
4. Make the result ATS-friendly using standard headings and plain text.
5. Use strong action verbs where supported by the original content.
6. Prioritize measurable achievements already present in the source.
7. Return a complete resume, not commentary about the resume.
8. Do not use tables, columns, emojis or decorative symbols.
9. Keep personal information appropriate for the target country.
10. Do not add a photo requirement.
11. Preserve the person's real identity and contact information if present.

SOURCE RESUME:
---BEGIN RESUME---
{resume}
---END RESUME---

Return the tailored resume with these sections where the source supports them:
NAME / CONTACT
PROFESSIONAL SUMMARY
CORE SKILLS
PROFESSIONAL EXPERIENCE
PROJECTS
EDUCATION
CERTIFICATIONS
ADDITIONAL INFORMATION
"""
    r = requests.post(
        f"{OLLAMA_URL.rstrip('/')}/api/generate",
        json={
            "model": OLLAMA_MODEL,
            "prompt": prompt,
            "stream": False,
            "options": {"temperature": 0.2}
        },
        timeout=180
    )
    r.raise_for_status()
    data = r.json()
    result = data.get("response", "").strip()
    if not result:
        raise RuntimeError("Ollama returned an empty response")
    return result

def make_docx(text, path):
    doc = Document()
    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        first = lines[0].strip()
        if first.upper() in {
            "PROFESSIONAL SUMMARY", "CORE SKILLS", "PROFESSIONAL EXPERIENCE",
            "PROJECTS", "EDUCATION", "CERTIFICATIONS", "ADDITIONAL INFORMATION",
            "TARGET ROLE", "TARGET COUNTRY", "NAME / CONTACT"
        }:
            p = doc.add_paragraph()
            run = p.add_run(first)
            run.bold = True
            for line in lines[1:]:
                doc.add_paragraph(line)
        else:
            for line in lines:
                doc.add_paragraph(line)
    doc.save(path)

def make_pdf(text, path):
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontName = "Helvetica"
    body.fontSize = 9.5
    body.leading = 13
    doc = SimpleDocTemplate(path, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    story = []
    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        first = lines[0].strip()
        if first.upper() in {
            "PROFESSIONAL SUMMARY", "CORE SKILLS", "PROFESSIONAL EXPERIENCE",
            "PROJECTS", "EDUCATION", "CERTIFICATIONS", "ADDITIONAL INFORMATION",
            "TARGET ROLE", "TARGET COUNTRY", "NAME / CONTACT"
        }:
            story.append(Paragraph(f"<b>{first}</b>", body))
            story.append(Spacer(1, 5))
            content = "<br/>".join(escape_html(x) for x in lines[1:])
        else:
            content = "<br/>".join(escape_html(x) for x in lines)
        if content:
            story.append(Paragraph(content, body))
            story.append(Spacer(1, 8))
    doc.build(story)

def escape_html(s):
    return (
        s.replace("&", "&amp;")
         .replace("<", "&lt;")
         .replace(">", "&gt;")
    )

@app.errorhandler(RequestEntityTooLarge)
def file_too_large(_error):
    if request.path.startswith("/api/"):
        return jsonify({"error": "File is too large. Maximum size is 8 MB."}), 413
    return render_template("page.html", title="File too large", kicker="Upload", paragraphs=["Maximum file size is 8 MB."]), 413

@app.errorhandler(404)
def not_found(_error):
    if request.path.startswith("/api/"):
        return jsonify({"error": "Not found."}), 404
    return render_template("page.html", title="Page not found", kicker="404", paragraphs=["That page is not on this site."]), 404

def featured_countries():
    return [BY_NAME[name] for name in FEATURED]

@app.route("/")
def index():
    return render_template("index.html", featured=featured_countries(), country_count=len(COUNTRIES))

@app.get("/tailor")
def tailor_page():
    selected = request.args.get("country", "")
    if selected not in COUNTRY_GUIDANCE:
        selected = ""
    return render_template("tailor.html", countries=sorted(COUNTRIES, key=lambda c: c["name"]), selected=selected)

@app.get("/countries")
def countries_page():
    return render_template("countries.html", countries=sorted(COUNTRIES, key=lambda c: c["name"]))

@app.get("/countries/<slug>")
def country_page(slug):
    country = BY_SLUG.get(slug)
    if not country:
        abort(404)
    return render_template("country.html", country=country)

@app.get("/blog")
def blog_page():
    return render_template("blog.html", posts=BLOG_POSTS)

@app.get("/blog/<slug>")
def blog_post(slug):
    post = POSTS_BY_SLUG.get(slug)
    if not post:
        abort(404)
    return render_template("post.html", post=post)

@app.get("/how-it-works")
def how_page():
    return render_template("how.html")

@app.get("/pricing")
def pricing_page():
    return render_template("pricing.html")

@app.get("/faq")
def faq_page():
    return render_template("faq.html", faqs=FAQS)

@app.get("/about")
def about_page():
    return render_template("page.html", title="About Jobaround", kicker="Company", paragraphs=[
        "Jobaround rewrites a resume for the country you are applying in. It changes structure, tone, and which personal details to leave off. It does not invent employers, dates, or skills.",
        "The site is a Flask app. When an Ollama server is configured, that model writes the draft. If it is unreachable, you still get a PDF and DOCX built from your original text.",
    ])

@app.get("/careers")
def careers_page():
    return render_template("page.html", title="Careers", kicker="Company", paragraphs=[
        "There are no open roles right now.",
        "If you want to talk about the project, use the contact form.",
    ])

@app.get("/press")
def press_page():
    return render_template("page.html", title="Press", kicker="Company", paragraphs=[
        "For a product question or a press note, use the contact form.",
    ])

@app.get("/privacy")
def privacy_page():
    return render_template("page.html", title="Privacy Policy", kicker="Legal", paragraphs=[
        "The uploaded resume is read for text and then deleted from the upload folder. The generated PDF and DOCX remain on the server so your download links keep working.",
        "If you are signed in, Jobaround stores your name, email, a password hash, and a short preview of each tailored resume in a local database.",
        "Contact messages are stored the same way. This app does not sell that information. A session cookie is used only to keep you signed in.",
    ])

@app.get("/terms")
def terms_page():
    return render_template("page.html", title="Terms of Use", kicker="Legal", paragraphs=[
        "Upload only resumes you have the right to process. You are responsible for checking the result before you send it to an employer.",
        "Jobaround does not promise interviews, offers, or that a generated file meets every local legal requirement.",
        "The service is free to use. A country note is guidance, not legal advice.",
    ])

@app.get("/cookies")
def cookies_page():
    return render_template("page.html", title="Cookie Policy", kicker="Legal", paragraphs=[
        "Jobaround sets one session cookie so sign-in and the upload form can stay on the same visit.",
        "The cookie is not used for advertising. Clearing it signs you out.",
    ])

@app.route("/contact", methods=["GET", "POST"])
def contact_page():
    error = None
    sent = False
    if request.method == "POST":
        if not csrf_ok():
            error = "Refresh the page and try again."
        else:
            name = request.form.get("name", "").strip()[:80]
            email = request.form.get("email", "").strip().lower()[:120]
            body = request.form.get("body", "").strip()[:4000]
            if not name or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email) or len(body) < 10:
                error = "Add your name, a valid email, and a message of at least 10 characters."
            else:
                conn = db()
                try:
                    conn.execute(
                        "INSERT INTO messages (name, email, body, created_at) VALUES (?, ?, ?, ?)",
                        (name, email, body, now()),
                    )
                    conn.commit()
                    sent = True
                finally:
                    conn.close()
    return render_template("contact.html", error=error, sent=sent)

@app.route("/signin", methods=["GET", "POST"])
def signin_page():
    if current_user():
        return redirect(url_for("account_page"))
    return auth_form("signin")

@app.route("/signup", methods=["GET", "POST"])
def signup_page():
    if current_user():
        return redirect(url_for("account_page"))
    return auth_form("signup")

def auth_form(mode):
    error = None
    if request.method == "POST":
        if not csrf_ok():
            error = "Refresh the page and try again."
        else:
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            conn = db()
            try:
                if mode == "signup":
                    name = request.form.get("name", "").strip()[:80]
                    if len(name) < 2 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email) or len(password) < 8:
                        error = "Use your name, a valid email, and a password of at least 8 characters."
                    else:
                        try:
                            conn.execute(
                                "INSERT INTO users (name, email, password_hash, created_at) VALUES (?, ?, ?, ?)",
                                (name, email, generate_password_hash(password), now()),
                            )
                            conn.commit()
                        except sqlite3.IntegrityError:
                            error = "An account with that email already exists."
                        else:
                            row = conn.execute("SELECT id FROM users WHERE email = ?", (email,)).fetchone()
                            session["user_id"] = row["id"]
                            return redirect(url_for("account_page"))
                else:
                    row = conn.execute("SELECT id, password_hash FROM users WHERE email = ?", (email,)).fetchone()
                    if not row or not check_password_hash(row["password_hash"], password):
                        error = "Email or password is incorrect."
                    else:
                        session["user_id"] = row["id"]
                        return redirect(url_for("account_page"))
            finally:
                conn.close()
    return render_template("auth.html", mode=mode, error=error)

@app.post("/logout")
def logout():
    if csrf_ok():
        session.pop("user_id", None)
    return redirect(url_for("index"))

@app.get("/account")
def account_page():
    user = current_user()
    if not user:
        return redirect(url_for("signin_page"))
    conn = db()
    try:
        rows = conn.execute(
            "SELECT country, job_title, preview, pdf_name, docx_name, ai_used, created_at FROM resumes WHERE user_id = ? ORDER BY id DESC LIMIT 20",
            (user["id"],),
        ).fetchall()
    finally:
        conn.close()
    history = []
    for row in rows:
        item = dict(row)
        item["flag"] = BY_NAME.get(item["country"], {}).get("flag", "")
        item["pdf_ok"] = os.path.exists(os.path.join(OUTPUT_DIR, item["pdf_name"] or ""))
        item["docx_ok"] = os.path.exists(os.path.join(OUTPUT_DIR, item["docx_name"] or ""))
        history.append(item)
    return render_template("account.html", history=history)

@app.post("/api/tailor")
def tailor():
    if not csrf_ok():
        return jsonify({"error": "Your session expired. Refresh the page and try again."}), 400
    uploaded = request.files.get("resume")
    country = request.form.get("country", "").strip()
    job_title = request.form.get("job_title", "").strip()
    job_description = request.form.get("job_description", "").strip()

    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Please upload a resume."}), 400
    if country not in COUNTRY_GUIDANCE:
        return jsonify({"error": "Please select a supported country."}), 400

    ext = Path(uploaded.filename).suffix.lower().lstrip(".")
    if ext not in ALLOWED:
        return jsonify({"error": "Upload PDF, DOCX or TXT only."}), 400

    safe = secure_filename(uploaded.filename) or "resume.txt"
    input_path = os.path.join(UPLOAD_DIR, f"{secrets.token_hex(4)}_{safe}")

    try:
        uploaded.save(input_path)
        resume = clean_text(extract_text(input_path))
        if len(resume) < 50:
            return jsonify({"error": "Could not extract enough text from the resume. Try a text-based PDF or DOCX."}), 400

        used_ai = True
        try:
            tailored = tailor_with_ollama(resume, country, job_title, job_description)
        except Exception:
            used_ai = False
            tailored = fallback_tailor(resume, country, job_title, job_description)

        stem = f"{secrets.token_hex(3)}_{Path(safe).stem}"[:48]
        docx_name = f"{stem}_jobaround.docx"
        pdf_name = f"{stem}_jobaround.pdf"
        docx_path = os.path.join(OUTPUT_DIR, docx_name)
        pdf_path = os.path.join(OUTPUT_DIR, pdf_name)
        make_docx(tailored, docx_path)
        make_pdf(tailored, pdf_path)
        if session.get("user_id"):
            conn = db()
            try:
                conn.execute(
                    "INSERT INTO resumes (user_id, country, job_title, preview, pdf_name, docx_name, ai_used, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (session["user_id"], country, job_title, tailored[:4000], pdf_name, docx_name, int(used_ai), now()),
                )
                conn.commit()
            finally:
                conn.close()

        return jsonify({
            "success": True,
            "ai_used": used_ai,
            "message": "Resume tailored successfully." if used_ai else "Resume processed using fallback mode because the local AI service was unavailable.",
            "preview": tailored,
            "docx": f"/download/docx/{docx_name}",
            "pdf": f"/download/pdf/{pdf_name}"
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        try:
            os.remove(input_path)
        except OSError:
            pass

@app.get("/download/<kind>/<filename>")
def download(kind, filename):
    if kind not in {"docx", "pdf"}:
        return "Invalid file type", 400
    safe = secure_filename(filename)
    path = os.path.join(OUTPUT_DIR, safe)
    if not os.path.exists(path):
        return "File not found", 404
    return send_file(path, as_attachment=True)

@app.get("/health")
def health():
    return jsonify({"status": "ok", "ollama_url": OLLAMA_URL, "ollama_model": OLLAMA_MODEL})

if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=True)
